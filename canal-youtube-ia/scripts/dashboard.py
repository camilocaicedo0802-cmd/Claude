"""Genera el dashboard de seguimiento (fase 12): dashboard_canal.xlsx.

Hojas: Leeme, Videos, Ideas, Experimentos, Resumen, Referencias, Glosario.
Las columnas calculadas usan fórmulas (se recalculan al escribir datos nuevos).

Uso: python3 scripts/dashboard.py  ->  dashboard_canal.xlsx
Después: recalcular con LibreOffice para guardar los valores en caché.
"""
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "dashboard_canal.xlsx"
N = 400  # filas preparadas en Videos

FUENTE = "Arial"
AZUL = Font(name=FUENTE, color="0000FF")  # entradas
NEGRO = Font(name=FUENTE, color="000000")  # fórmulas
VERDE = Font(name=FUENTE, color="008000")  # vínculos a otra hoja
NEGRITA = Font(name=FUENTE, bold=True)
TITULO = Font(name=FUENTE, bold=True, size=14)
CABECERA = Font(name=FUENTE, bold=True, color="FFFFFF")
FONDO_CAB = PatternFill("solid", fgColor="1E1A16")
FONDO_CALC = PatternFill("solid", fgColor="EADBC0")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
FONDO_GANA = PatternFill("solid", fgColor="C6EFCE")
FONDO_PIERDE = PatternFill("solid", fgColor="FFC7CE")
BORDE = Border(bottom=Side(style="thin", color="BBBBBB"))
AJUSTE = Alignment(wrap_text=True, vertical="top")


def cabecera(ws, fila, textos, calc_desde=None):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = CABECERA
        c.fill = FONDO_CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if calc_desde and i >= calc_desde:
            c.fill = PatternFill("solid", fgColor="5A4632")
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def anchos(ws, valores):
    for i, w in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------- Videos
COLS_VIDEOS = [
    ("id", 7), ("estado_video", 12), ("fecha_publicacion", 12), ("formato", 9), ("pilar", 7),
    ("tema", 30), ("keyword_principal", 24), ("titulo", 40), ("concepto_miniatura", 30), ("url", 28),
    ("duracion_min", 9), ("vistas_24h", 10), ("vistas_7d", 10), ("vistas_28d", 10),
    ("impresiones_28d", 11), ("ctr_28d", 8), ("avd_min", 8), ("pct_visto", 9), ("retencion_30s", 10),
    ("horas_vistas_28d", 10), ("subs_ganados_28d", 10), ("trafico_busqueda", 10),
    ("trafico_recomendaciones", 13), ("trafico_shorts", 10), ("trafico_otros", 10),
    ("costo_usd", 9), ("horas_produccion", 10), ("estado_monetizacion", 18), ("cumplimiento_ok", 11),
    ("aprendizaje", 36), ("proximo_experimento", 30),
    ("n_formato (aux)", 8), ("mediana_prev5_7d", 11), ("multiplo_7d", 10), ("mediana_prev5_28d", 11),
    ("multiplo_28d", 10), ("mediana_pct_visto_formato", 12), ("clasificacion", 22),
    ("costo_por_mil_vistas", 11), ("subs_por_mil_vistas", 11),
]
IDX = {nombre: i + 1 for i, (nombre, _) in enumerate(COLS_VIDEOS)}
L = {nombre: get_column_letter(i) for nombre, i in IDX.items()}
PRIMERA_CALC = IDX["n_formato (aux)"]

# Plan de la fase 11 (fechas previstas; estado "planificado")
PLAN = [
    ("L1", date(2026, 11, 19), "Largo", "P1", "La verdadera historia de La Llorona", "la verdadera historia de la llorona",
     "La verdadera historia de La Llorona (según las crónicas)", "A: ¿EXISTIÓ? / B: ANTES DE 1519 / C: LO DICE AQUÍ"),
    ("S1", date(2026, 11, 21), "Short", "P1", "Los mexicas oyeron a La Llorona antes de la conquista (según Sahagún)", "la llorona", "", ""),
    ("S2", date(2026, 11, 26), "Short", "P1", "¿Cómo se llama La Llorona en tu país?", "la llorona", "", ""),
    ("L2", date(2026, 12, 3), "Largo", "P2", "¿Por qué Colombia se llama Colombia?", "historia de colombia",
     "¿Por qué Colombia se llama Colombia?", "A: ¿POR QUÉ ÉL? / B: ¿Y ANTES? / C: ERA TODO"),
    ("S3", date(2026, 12, 5), "Short", "P2", "Colombia se llamó Estados Unidos de Colombia", "historia de colombia", "", ""),
    ("S4", date(2026, 12, 10), "Short", "P2", "El nombre 'Colombia' era para todo el continente", "historia de colombia", "", ""),
    ("L3", date(2026, 12, 17), "Largo", "P1", "El Dorado: lo que sí existió", "el dorado leyenda",
     "El Dorado sí existió (pero no era una ciudad)", "A: NO ERA CIUDAD / B: ESTO SÍ / C: ¿DÓNDE ESTÁ?"),
    ("S5", date(2026, 12, 19), "Short", "P1", "El Dorado no era una ciudad. Era un hombre", "el dorado", "", ""),
    ("S6", date(2026, 12, 24), "Short", "P1", "Intentaron vaciar una laguna para sacar el oro", "laguna de guatavita", "", ""),
    ("L4", date(2026, 12, 30), "Largo", "P1", "El chupacabras: cómo nació un monstruo moderno", "la leyenda del chupacabras",
     "El chupacabras: cómo nació un monstruo en 1995", "A: NACIÓ EN 1995 / B: ¿REAL? / C: LO CREÓ LA TV"),
    ("S7", date(2027, 1, 2), "Short", "P1", "El chupacabras tiene fecha de nacimiento", "chupacabras", "", ""),
    ("S8", date(2027, 1, 6), "Short", "P1", "Quién le puso el nombre al chupacabras (verificar)", "chupacabras", "", ""),
    ("S9", date(2027, 1, 9), "Short", "P2", "Colombia lleva el nombre de alguien que nunca vino (solo si se verifica)", "historia de colombia", "", ""),
    ("L5", date(2027, 1, 14), "Largo", "P2", "¿De verdad unos cientos de españoles conquistaron a los mexicas?", "conquista de mexico",
     "¿De verdad 500 españoles conquistaron a los mexicas?", "A: ¿SOLO 500? / B: NO ESTABA SOLO / C: TODOS CONTRA…"),
    ("S10", date(2027, 1, 16), "Short", "P2", "Cortés no conquistó México solo", "hernan cortes", "", ""),
    ("S11", date(2027, 1, 20), "Short", "P2", "El arma más letal de la conquista no era una espada", "conquista de mexico", "", ""),
    ("S12", date(2027, 1, 23), "Short", "P2", "Ciudad de México está construida sobre un lago", "historia de mexico", "", ""),
    ("L6", date(2027, 1, 28), "Largo", "P2", "¿Por qué la Ciudad de México está sobre un lago?", "historia de mexico",
     "¿Por qué la Ciudad de México está sobre un lago?", "A: AQUÍ HABÍA AGUA / B: SE HUNDE / C: ¿POR QUÉ AQUÍ?"),
    ("S13", date(2027, 1, 30), "Short", "P2", "La ciudad se hunde cada año (con fuente oficial)", "cdmx", "", ""),
]


def hoja_videos(wb):
    ws = wb.create_sheet("Videos")
    cabecera(ws, 1, [n for n, _ in COLS_VIDEOS], calc_desde=PRIMERA_CALC)
    anchos(ws, [w for _, w in COLS_VIDEOS])
    ws.row_dimensions[1].height = 42

    entradas = [n for n, _ in COLS_VIDEOS[:PRIMERA_CALC - 1] if n not in ("pct_visto", "trafico_otros")]
    for fila in range(2, N + 2):
        for nombre in entradas:
            ws.cell(row=fila, column=IDX[nombre]).font = AZUL
        r = fila
        D, K, Q = L["formato"], L["duracion_min"], L["avd_min"]
        ws[f"{L['pct_visto']}{r}"] = f'=IF(AND(ISNUMBER({K}{r}),ISNUMBER({Q}{r}),{K}{r}>0),{Q}{r}/{K}{r},"")'
        V, X = L["trafico_busqueda"], L["trafico_shorts"]
        ws[f"{L['trafico_otros']}{r}"] = f'=IF(COUNT({V}{r}:{X}{r})=0,"",MAX(0,1-SUM({V}{r}:{X}{r})))'
        AF = L["n_formato (aux)"]
        ws[f"{AF}{r}"] = f'=IF({D}{r}="","",COUNTIF(${D}$2:{D}{r},{D}{r}))'
        for vistas, med, mult in (("vistas_7d", "mediana_prev5_7d", "multiplo_7d"),
                                  ("vistas_28d", "mediana_prev5_28d", "multiplo_28d")):
            Vc = L[vistas]
            ref = f"{L[med]}{r}"
            ws[ref] = ArrayFormula(
                ref,
                f'=IF({D}{r}="","",IFERROR(MEDIAN(IF((${D}$2:${D}${N + 1}={D}{r})'
                f'*(${AF}$2:${AF}${N + 1}>={AF}{r}-5)*(${AF}$2:${AF}${N + 1}<={AF}{r}-1)'
                f'*(${Vc}$2:${Vc}${N + 1}<>""),${Vc}$2:${Vc}${N + 1})),""))',
            )
            ws[f"{L[mult]}{r}"] = (f'=IF(AND(ISNUMBER({Vc}{r}),ISNUMBER({L[med]}{r})),'
                                   f'IF({L[med]}{r}>0,{Vc}{r}/{L[med]}{r},""),"")')
        R = L["pct_visto"]
        ref = f"{L['mediana_pct_visto_formato']}{r}"
        ws[ref] = ArrayFormula(
            ref,
            f'=IF({D}{r}="","",IFERROR(MEDIAN(IF((${D}$2:${D}${N + 1}={D}{r})'
            f'*ISNUMBER(${R}$2:${R}${N + 1}),${R}$2:${R}${N + 1})),""))',
        )
        m7, m28, mp = L["multiplo_7d"], L["multiplo_28d"], L["mediana_pct_visto_formato"]
        ws[f"{L['clasificacion']}{r}"] = (
            f'=IF(ISNUMBER({m28}{r}),'
            f'IF(AND({m28}{r}>=Resumen!$C$5,OR(NOT(ISNUMBER({R}{r})),NOT(ISNUMBER({mp}{r})),{R}{r}>={mp}{r})),"Ganador",'
            f'IF({m28}{r}<=Resumen!$C$6,"Perdedor","Normal")),'
            f'IF(ISNUMBER({m7}{r}),'
            f'IF({m7}{r}>=Resumen!$C$5,"Ganador (7 d, provisional)",'
            f'IF({m7}{r}<=Resumen!$C$6,"Perdedor (7 d, provisional)","Normal (7 d, provisional)")),""))'
        )
        Z, N28, U = L["costo_usd"], L["vistas_28d"], L["subs_ganados_28d"]
        ws[f"{L['costo_por_mil_vistas']}{r}"] = f'=IF(AND(ISNUMBER({Z}{r}),ISNUMBER({N28}{r})),IF({N28}{r}>0,{Z}{r}/{N28}{r}*1000,""),"")'
        ws[f"{L['subs_por_mil_vistas']}{r}"] = f'=IF(AND(ISNUMBER({U}{r}),ISNUMBER({N28}{r})),IF({N28}{r}>0,{U}{r}/{N28}{r}*1000,""),"")'
        for nombre, _ in COLS_VIDEOS[PRIMERA_CALC - 1:]:
            ws.cell(row=r, column=IDX[nombre]).font = NEGRO
        for nombre in ("pct_visto", "trafico_otros"):
            ws.cell(row=r, column=IDX[nombre]).font = NEGRO

    for i, (id_, fecha, fmt, pilar, tema, kw, titulo, mini) in enumerate(PLAN, start=2):
        valores = {"id": id_, "estado_video": "planificado", "fecha_publicacion": fecha, "formato": fmt,
                   "pilar": pilar, "tema": tema, "keyword_principal": kw, "titulo": titulo or None,
                   "concepto_miniatura": mini or None, "cumplimiento_ok": None}
        for k, v in valores.items():
            if v is not None:
                ws.cell(row=i, column=IDX[k], value=v)

    # Formatos de número
    for r in range(2, N + 2):
        ws[f"{L['fecha_publicacion']}{r}"].number_format = "yyyy-mm-dd"
        for n in ("ctr_28d", "pct_visto", "retencion_30s", "trafico_busqueda", "trafico_recomendaciones",
                  "trafico_shorts", "trafico_otros", "mediana_pct_visto_formato"):
            ws[f"{L[n]}{r}"].number_format = "0.0%"
        for n in ("vistas_24h", "vistas_7d", "vistas_28d", "impresiones_28d", "subs_ganados_28d",
                  "mediana_prev5_7d", "mediana_prev5_28d"):
            ws[f"{L[n]}{r}"].number_format = "#,##0"
        for n in ("multiplo_7d", "multiplo_28d"):
            ws[f"{L[n]}{r}"].number_format = '0.00"×"'
        for n in ("duracion_min", "avd_min", "horas_vistas_28d", "horas_produccion", "subs_por_mil_vistas"):
            ws[f"{L[n]}{r}"].number_format = "0.0"
        for n in ("costo_usd", "costo_por_mil_vistas"):
            ws[f"{L[n]}{r}"].number_format = '"$"#,##0.00'

    listas = [
        ("estado_video", '"idea,planificado,en producción,programado,publicado"'),
        ("formato", '"Largo,Short"'),
        ("pilar", '"P1,P2,P3,P4"'),
        ("estado_monetizacion", '"sin YPP,YPP ampliado,YPP anuncios - completos,YPP anuncios - limitados,sin anuncios"'),
        ("cumplimiento_ok", '"sí,no"'),
    ]
    for nombre, formula in listas:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{L[nombre]}2:{L[nombre]}{N + 1}")

    rango = f"{L['clasificacion']}2:{L['clasificacion']}{N + 1}"
    col = L["clasificacion"]
    ws.conditional_formatting.add(rango, FormulaRule(formula=[f'LEFT({col}2,7)="Ganador"'], fill=FONDO_GANA))
    ws.conditional_formatting.add(rango, FormulaRule(formula=[f'LEFT({col}2,8)="Perdedor"'], fill=FONDO_PIERDE))

    notas = {
        "vistas_24h": "YouTube Studio > Contenido > video > Estadísticas, 24 h después de publicar.",
        "ctr_28d": "Porcentaje de clics de las impresiones (Studio > Alcance), a los 28 días. Escribe 0,045 para 4,5 %.",
        "avd_min": "Duración media de visualización en minutos (Studio > Interacción). 3:30 = 3,5.",
        "retencion_30s": "Porcentaje de espectadores que siguen viendo en el segundo 30 (curva de retención).",
        "horas_vistas_28d": "Tiempo de visualización en horas (Studio). Solo las de largos cuentan para el Programa de Socios.",
        "trafico_busqueda": "Fuentes de tráfico (Studio > Alcance): porcentaje de vistas desde la búsqueda de YouTube.",
        "trafico_recomendaciones": "Funciones de exploración + videos sugeridos.",
        "costo_usd": "Parte del costo mensual de herramientas imputada a este video (p. ej., 41 USD / videos del mes).",
        "cumplimiento_ok": "Lista de comprobación de la fase 10, sección 14, completa.",
        "mediana_prev5_7d": "Mediana de vistas a 7 días de los 5 videos anteriores del mismo formato.",
        "multiplo_28d": "Vistas a 28 días ÷ mediana de los 5 anteriores del mismo formato. Indicador principal.",
        "clasificacion": "Ganador: múltiplo 28 d ≥ umbral y % visto ≥ mediana del formato. Perdedor: múltiplo ≤ umbral inferior. Umbrales en Resumen.",
    }
    for nombre, texto in notas.items():
        ws.cell(row=1, column=IDX[nombre]).comment = Comment(texto, "dashboard")
    return ws


# ---------------------------------------------------------------- Ideas
PESOS = [("demanda", 20), ("curiosidad", 15), ("encaje_promesa", 15), ("evidencia_fuentes", 15),
         ("diferenciacion", 10), ("apta_anunciantes", 10), ("shorts_serie", 5), ("costo_bajo", 10)]


def hoja_ideas(wb):
    ws = wb.create_sheet("Ideas")
    cols = ["id", "pilar", "tema", "keyword_principal", "evidencia_demanda"] + [p for p, _ in PESOS] + \
           ["total_100", "decision", "estado", "video_id"]
    cabecera(ws, 1, cols)
    anchos(ws, [6, 6, 44, 30, 60] + [10] * 8 + [10, 22, 14, 9])
    ws.row_dimensions[1].height = 30
    ws["A2"] = "PESOS"
    ws["A2"].font = NEGRITA
    ws["E2"] = "Pesos de la fase 8 (suman 100). Puedes cambiarlos: los totales se recalculan."
    ws["E2"].font = Font(name=FUENTE, italic=True)
    for j, (_, w) in enumerate(PESOS, start=6):
        c = ws.cell(row=2, column=j, value=w)
        c.font = AZUL
        c.fill = AMARILLO
    ws.cell(row=2, column=14, value="=SUM(F2:M2)").font = NEGRO
    with (RAIZ / "datos" / "banco_ideas.csv").open(encoding="utf-8") as fh:
        filas = list(csv.DictReader(fh))
    for i, f in enumerate(filas, start=3):
        for j, k in enumerate(["id", "pilar", "tema", "keyword_principal", "evidencia_demanda"], start=1):
            ws.cell(row=i, column=j, value=f[k]).font = AZUL
        for j, (p, _) in enumerate(PESOS, start=6):
            ws.cell(row=i, column=j, value=int(f[p])).font = AZUL
        ws.cell(row=i, column=14, value=f"=SUMPRODUCT(F{i}:M{i},$F$2:$M$2)/5").font = NEGRO
        ws.cell(row=i, column=15,
                value=f'=IF(N{i}>=75,"producir pronto",IF(N{i}>=60,"banco","reformular o descartar"))').font = NEGRO
        estado = "planificado" if f["id"] in ("I01", "I02", "I04", "I05", "I07", "I09") else "pendiente"
        ws.cell(row=i, column=16, value=estado).font = AZUL
        vid = {"I01": "L1", "I02": "L2", "I04": "L3", "I07": "L4", "I05": "L5", "I09": "L6"}.get(f["id"])
        if vid:
            ws.cell(row=i, column=17, value=vid).font = AZUL
        for j in range(1, 18):
            ws.cell(row=i, column=j).alignment = AJUSTE
    ultima = 2 + len(filas)
    for r in range(ultima + 1, ultima + 41):  # filas vacías preparadas
        ws.cell(row=r, column=14, value=f'=IF(COUNT(F{r}:M{r})=0,"",SUMPRODUCT(F{r}:M{r},$F$2:$M$2)/5)').font = NEGRO
        ws.cell(row=r, column=15,
                value=f'=IF(N{r}="","",IF(N{r}>=75,"producir pronto",IF(N{r}>=60,"banco","reformular o descartar")))').font = NEGRO
        for j in list(range(1, 14)) + [16, 17]:
            ws.cell(row=r, column=j).font = AZUL
    for r in range(3, ultima + 41):
        ws.cell(row=r, column=14).number_format = "0.0"
    ws.cell(row=1, column=14).comment = Comment("Σ(puntos × peso) ÷ 5. Escala de 1 a 5 por criterio (fase 8).", "dashboard")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F3:M{ultima + 40}")
    return ws


# ---------------------------------------------------------------- Experimentos
def hoja_experimentos(wb):
    ws = wb.create_sheet("Experimentos")
    cols = ["id", "fecha_inicio", "video_id", "tipo", "regla_que_se_prueba", "hipotesis", "variante_A",
            "variante_B", "variante_C", "metrica_de_decision", "resultado", "dato_resultado", "conclusion",
            "victorias_de_la_regla", "estado_de_la_regla"]
    cabecera(ws, 1, cols, calc_desde=14)
    anchos(ws, [6, 12, 9, 18, 26, 50, 22, 22, 22, 24, 16, 22, 40, 11, 14])
    ws.row_dimensions[1].height = 30
    filas = [
        ("H1", date(2026, 11, 19), "L1", "A/B miniatura", "pregunta en miniatura",
         "La miniatura con pregunta (¿EXISTIÓ?) gana a la de afirmación (ANTES DE 1519) en tiempo de visualización",
         "¿EXISTIÓ?", "ANTES DE 1519", "LO DICE AQUÍ", "Resultado de Prueba y compara (tiempo de visualización)"),
        ("H2", date(2026, 12, 3), "L2", "Pilar", "P2 trae más búsqueda",
         "Un '¿Por qué…?' de identidad nacional (P2) consigue más tráfico de búsqueda que una leyenda (P1)",
         "L2 (P2)", "L1 (P1)", "", "trafico_busqueda a 7 días"),
        ("H3", date(2026, 12, 17), "L3", "A/B título", "título contraintuitivo",
         "El título contraintuitivo ('sí existió, pero no era una ciudad') gana a 'Lo que es verdad de El Dorado'",
         "El Dorado sí existió (pero no era una ciudad)", "Lo que es verdad de El Dorado", "", "Resultado de Prueba y compara"),
        ("H4", date(2026, 12, 30), "L4", "Tema", "leyenda moderna con fecha",
         "Una leyenda moderna con fecha (1995) retiene más que una antigua (La Llorona)",
         "L4", "L1", "", "avd_min y pct_visto a 28 días"),
        ("H5", date(2027, 1, 14), "L5", "Tema", "debate sin polarización",
         "Un tema de debate tiene más comentarios sin pasar a la polarización",
         "L5", "mediana de largos anteriores", "", "Comentarios por 1.000 vistas y comentarios moderados"),
    ]
    for i, f in enumerate(filas, start=2):
        for j, v in enumerate(f, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = AZUL
            c.alignment = AJUSTE
    for r in range(2, 202):
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=14,
                value=f'=IF(E{r}="","",COUNTIFS($E$2:$E$201,E{r},$K$2:$K$201,"Winner"))').font = NEGRO
        ws.cell(row=r, column=15,
                value=f'=IF(E{r}="","",IF(N{r}>=3,"Norma del canal",IF(N{r}>=1,"Hipótesis con 1-2 victorias","Hipótesis")))').font = NEGRO
        for j in range(1, 14):
            if ws.cell(row=r, column=j).value is None:
                ws.cell(row=r, column=j).font = AZUL
    dv = DataValidation(type="list", formula1='"Winner,Performed Same,Inconclusive,Sin resultado"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("K2:K201")
    ws.cell(row=1, column=15).comment = Comment(
        "Una regla pasa a ser norma cuando gana en 3 pruebas de temas distintos (fase 8, sección 6).", "dashboard")
    return ws


# ---------------------------------------------------------------- Resumen
def hoja_resumen(wb):
    ws = wb.create_sheet("Resumen")
    anchos(ws, [3, 46, 16, 16, 16, 16, 40])
    ws["B1"] = "Resumen del canal"
    ws["B1"].font = TITULO
    ws["B2"] = "Celdas amarillas: las rellenas tú con datos de YouTube Studio. El resto se calcula solo."
    ws["B2"].font = Font(name=FUENTE, italic=True)

    ws["B4"] = "Parámetros"
    ws["B4"].font = NEGRITA
    params = [
        (5, "Umbral de ganador (múltiplo a 28 días ≥)", 1.5, "Fase 12. Mismo criterio de rendimiento relativo de la fase 1."),
        (6, "Umbral de perdedor (múltiplo a 28 días ≤)", 0.67, "Inverso de 1,5."),
    ]
    for fila, txt, val, nota in params:
        ws.cell(row=fila, column=2, value=txt).font = NEGRO
        c = ws.cell(row=fila, column=3, value=val)
        c.font = AZUL
        c.fill = AMARILLO
        c.number_format = '0.00"×"'
        ws.cell(row=fila, column=7, value=nota).font = Font(name=FUENTE, italic=True, color="666666")

    V = "Videos!"
    rng = lambda n: f"{V}${L[n]}$2:${L[n]}${N + 1}"  # noqa: E731
    ws["B8"] = "Producción y alcance (suma de las cifras a 28 días de cada video)"
    ws["B8"].font = NEGRITA
    ws["C9"], ws["D9"] = "Largos", "Shorts"
    for c in ("C9", "D9"):
        ws[c].font = NEGRITA
    kpis = [
        ("Videos publicados (con URL)", lambda f: f'=COUNTIFS({rng("formato")},"{f}",{rng("url")},"<>")', "0"),
        ("Vistas a 28 días (suma)", lambda f: f'=SUMIFS({rng("vistas_28d")},{rng("formato")},"{f}")', "#,##0"),
        ("Horas vistas a 28 días (suma)", lambda f: f'=SUMIFS({rng("horas_vistas_28d")},{rng("formato")},"{f}")', "#,##0.0"),
        ("Suscriptores ganados a 28 días (suma)", lambda f: f'=SUMIFS({rng("subs_ganados_28d")},{rng("formato")},"{f}")', "#,##0"),
        ("Horas de producción medias por video",
         lambda f: f'=IFERROR(AVERAGEIFS({rng("horas_produccion")},{rng("formato")},"{f}"),"")', "0.0"),
        ("Costo imputado total (USD)", lambda f: f'=SUMIFS({rng("costo_usd")},{rng("formato")},"{f}")', '"$"#,##0.00'),
        ("Ganadores (incluye provisionales)",
         lambda f: f'=COUNTIFS({rng("formato")},"{f}",{rng("clasificacion")},"Ganador*")', "0"),
        ("Perdedores (incluye provisionales)",
         lambda f: f'=COUNTIFS({rng("formato")},"{f}",{rng("clasificacion")},"Perdedor*")', "0"),
    ]
    fila = 10
    for txt, fn, fmt in kpis:
        ws.cell(row=fila, column=2, value=txt).font = NEGRO
        for col, f in ((3, "Largo"), (4, "Short")):
            c = ws.cell(row=fila, column=col, value=fn(f))
            c.font = VERDE
            c.number_format = fmt
        fila += 1

    fila += 1
    ws.cell(row=fila, column=2, value="Medianas por formato (videos con dato)").font = NEGRITA
    fila += 1
    ws.cell(row=fila, column=3, value="Largos").font = NEGRITA
    ws.cell(row=fila, column=4, value="Shorts").font = NEGRITA
    fila += 1
    medianas = [("CTR a 28 días", "ctr_28d", "0.0%"), ("Duración media vista (min)", "avd_min", "0.0"),
                ("% visto", "pct_visto", "0.0%"), ("Retención a 30 s", "retencion_30s", "0.0%"),
                ("Vistas a 28 días", "vistas_28d", "#,##0"), ("Múltiplo a 28 días", "multiplo_28d", '0.00"×"')]
    for txt, n, fmt in medianas:
        ws.cell(row=fila, column=2, value=txt).font = NEGRO
        for col, f in ((3, "Largo"), (4, "Short")):
            ref = f"{get_column_letter(col)}{fila}"
            ws[ref] = ArrayFormula(
                ref, f'=IFERROR(MEDIAN(IF(({rng("formato")}="{f}")*ISNUMBER({rng(n)}),{rng(n)})),"")')
            ws[ref].font = VERDE
            ws[ref].number_format = fmt
        fila += 1

    fila += 1
    ws.cell(row=fila, column=2, value="Por pilar (solo largos)").font = NEGRITA
    fila += 1
    for j, t in enumerate(["Largos con vistas a 28 d", "Mediana del múltiplo a 28 d", "Mediana de vistas a 28 d",
                           "Mediana de % visto"], start=3):
        ws.cell(row=fila, column=j, value=t).font = NEGRITA
        ws.cell(row=fila, column=j).alignment = AJUSTE
    fila += 1
    cond = lambda p, n: f'({rng("formato")}="Largo")*({rng("pilar")}="{p}")*ISNUMBER({rng(n)})'  # noqa: E731
    for p, nombre in (("P1", "P1 Leyendas"), ("P2", "P2 ¿Por qué…?"), ("P3", "P3 Así vivías"), ("P4", "P4 De cada país")):
        ws.cell(row=fila, column=2, value=nombre).font = NEGRO
        ws.cell(row=fila, column=3, value=f'=SUMPRODUCT({cond(p, "vistas_28d")})').font = VERDE
        for col, n, fmt in ((4, "multiplo_28d", '0.00"×"'), (5, "vistas_28d", "#,##0"), (6, "pct_visto", "0.0%")):
            ref = f"{get_column_letter(col)}{fila}"
            ws[ref] = ArrayFormula(ref, f'=IFERROR(MEDIAN(IF({cond(p, n)},{rng(n)})),"")')
            ws[ref].font = VERDE
            ws[ref].number_format = fmt
        fila += 1
    ws.cell(row=fila, column=2,
            value="Regla de la fase 11: un pilar se prioriza (mezcla 2:1) si su mediana de múltiplos es ≥ 1,5 veces la del otro.").font = Font(
        name=FUENTE, italic=True, color="666666")

    fila += 2
    ws.cell(row=fila, column=2, value="Progreso hacia el Programa de Socios (datos de Studio)").font = NEGRITA
    fila += 1
    entradas = [("Suscriptores actuales", 0, "#,##0"), ("Horas de visualización válidas (últimos 365 días)", 0, "#,##0"),
                ("Vistas válidas de Shorts (últimos 90 días)", 0, "#,##0"), ("Subidas públicas válidas (últimos 90 días)", 0, "0")]
    ref_ent = {}
    for txt, val, fmt in entradas:
        ws.cell(row=fila, column=2, value=txt).font = NEGRO
        c = ws.cell(row=fila, column=3, value=val)
        c.font = AZUL
        c.fill = AMARILLO
        c.number_format = fmt
        ref_ent[txt] = f"$C${fila}"
        fila += 1
    ws.cell(row=fila - 4, column=7,
            value="Fuente: YouTube Studio > Ganar dinero (muestra el progreso oficial). Requisitos: fase 10.").font = Font(
        name=FUENTE, italic=True, color="666666")
    fila += 1
    for j, t in enumerate(["Suscriptores", "Horas (365 d)", "Shorts (90 d)", "Progreso (suscriptores × horas)", "Vigencia"], start=3):
        ws.cell(row=fila, column=j, value=t).font = NEGRITA
    fila += 1
    sub, hor, sho = ref_ent["Suscriptores actuales"], ref_ent["Horas de visualización válidas (últimos 365 días)"], \
        ref_ent["Vistas válidas de Shorts (últimos 90 días)"]
    niveles = [("Ampliado (membresías, Super Thanks)", 500, 3000, 3000000, "sin cambios en 2027"),
               ("Anuncios: hasta el 31-ene-2027", 1000, 4000, 10000000, "solicitudes hasta el 31-ene-2027"),
               ("Anuncios: desde el 1-feb-2027 (canales nuevos)", 1000, 8000, 20000000, "1-feb-2027 en adelante")]
    for txt, s, h, sh, vig in niveles:
        ws.cell(row=fila, column=2, value=txt).font = NEGRO
        for col, v, fmt in ((3, s, "#,##0"), (4, h, "#,##0"), (5, sh, "#,##0")):
            c = ws.cell(row=fila, column=col, value=v)
            c.font = AZUL
            c.number_format = fmt
        c = ws.cell(row=fila, column=6, value=f"=MIN(1,{sub}/C{fila})*MIN(1,{hor}/D{fila})")
        c.font = NEGRO
        c.number_format = "0.0%"
        ws.cell(row=fila, column=7, value=vig).font = NEGRO
        fila += 1
    ws.cell(row=fila, column=2,
            value="Progreso = (suscriptores ÷ meta, máx. 100 %) × (horas ÷ meta, máx. 100 %). Es 100 % solo si se cumplen ambas.").font = Font(
        name=FUENTE, italic=True, color="666666")

    fila += 2
    ws.cell(row=fila, column=2, value="Escenarios de la fase 11 al día 90 (supuestos, no predicciones)").font = NEGRITA
    fila += 1
    for j, t in enumerate(["Conservador", "Intermedio", "Optimista", "Real (Studio)"], start=3):
        ws.cell(row=fila, column=j, value=t).font = NEGRITA
    fila += 1
    esc = [("Vistas de largos acumuladas", 2500, 20000, 175000, None),
           ("Vistas de Shorts acumuladas", 10000, 60000, 450000, None),
           ("Horas de visualización válidas", 125, 1170, 11700, f"={hor}"),
           ("Suscriptores", 22, 150, 1240, f"={sub}")]
    for txt, a, b, c_, real in esc:
        ws.cell(row=fila, column=2, value=txt).font = NEGRO
        for col, v in ((3, a), (4, b), (5, c_)):
            cc = ws.cell(row=fila, column=col, value=v)
            cc.font = AZUL
            cc.number_format = "#,##0"
        cr = ws.cell(row=fila, column=6, value=real if real else 0)
        cr.number_format = "#,##0"
        if real:
            cr.font = NEGRO
        else:
            cr.font = AZUL
            cr.fill = AMARILLO
        fila += 1
    ws.cell(row=fila, column=2,
            value="Supuestos de cada escenario y advertencias: 13-fase11-plan-90-dias.md, sección 3. Suscriptores: punto medio del rango.").font = Font(
        name=FUENTE, italic=True, color="666666")
    return ws


# ---------------------------------------------------------------- Referencias
def hoja_referencias(wb):
    ws = wb.create_sheet("Referencias")
    cols = ["canal", "idioma", "creado", "suscriptores", "vistas_totales", "videos_totales", "subidas_semana_90d",
            "mediana_vistas_largos", "duracion_mediana_min", "dias_a_1k", "dias_a_10k", "pct_vistas_top10pct",
            "suscriptores_por_vista"]
    cabecera(ws, 1, cols, calc_desde=13)
    anchos(ws, [26, 8, 12, 13, 15, 10, 11, 14, 11, 9, 9, 11, 12])
    ws.row_dimensions[1].height = 30
    with (RAIZ / "datos" / "canales.csv").open(encoding="utf-8") as fh:
        filas = list(csv.DictReader(fh))
    num = lambda x: float(x) if x not in ("", None) else None  # noqa: E731
    for i, f in enumerate(filas, start=2):
        vals = [f["nombre"], f["idioma_original"], f["canal_creado"], num(f["suscriptores"]), num(f["vistas_totales"]),
                num(f["videos_totales"]), num(f["subidas_por_semana_90d"]), num(f["mediana_vistas_largos"]),
                num(f["duracion_mediana_largos_min"]), num(f["dias_a_1k"]), num(f["dias_a_10k"]),
                num(f["pct_vistas_top10pct_videos"])]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = AZUL
        ws.cell(row=i, column=12).value = (vals[11] / 100) if vals[11] is not None else None
        ws.cell(row=i, column=12).number_format = "0%"
        ws.cell(row=i, column=13, value=f"=IF(E{i}>0,D{i}/E{i},\"\")").font = NEGRO
        ws.cell(row=i, column=13).number_format = "0.00%"
        for j in (4, 5, 8):
            ws.cell(row=i, column=j).number_format = "#,##0"
    nota = len(filas) + 3
    ws.cell(row=nota, column=1, value=(
        "Fuente: datos/canales.csv (fase 1, consulta del 27-oct-2026): página 'Acerca de', yt-dlp y vidIQ. "
        "Días a 1k/10k: historial de vidIQ (vacío = no consultado). Squire Explains tenía 4 semanas: su ritmo real era "
        "de ~9 publicaciones por semana, no 2,8. Son canales elegidos por su éxito (sesgo de supervivencia).")).font = Font(
        name=FUENTE, italic=True, color="666666")
    return ws


# ---------------------------------------------------------------- Glosario
GLOSARIO = [
    ("vistas_24h / 7d / 28d", "Vistas del video a 1, 7 y 28 días de publicarse", "Pública (el número actual); por día solo en Studio", "Studio > Contenido > video > Estadísticas", "24 h, 7 d, 28 d"),
    ("impresiones_28d", "Veces que la miniatura se mostró en YouTube", "Privada", "Studio > Alcance", "28 d"),
    ("ctr_28d", "Porcentaje de impresiones que acabaron en clic", "Privada", "Studio > Alcance", "28 d"),
    ("avd_min", "Duración media de visualización", "Privada", "Studio > Interacción", "28 d"),
    ("pct_visto", "AVD ÷ duración del video (calculado)", "Privada (derivada)", "Fórmula", "28 d"),
    ("retencion_30s", "Porcentaje que sigue viendo en el segundo 30", "Privada", "Studio > Interacción > Retención del público", "7 d"),
    ("horas_vistas_28d", "Tiempo de visualización en horas", "Privada", "Studio > Interacción", "28 d"),
    ("subs_ganados_28d", "Suscriptores ganados desde ese video", "Privada", "Studio > Audiencia (filtrar por video)", "28 d"),
    ("trafico_*", "Reparto de vistas por fuente de tráfico", "Privada", "Studio > Alcance > Fuentes de tráfico", "28 d"),
    ("likes, comentarios", "Interacción visible", "Pública", "Página del video", "28 d (opcional)"),
    ("duracion_min, título, fecha", "Datos del video", "Pública", "Página del video", "Al publicar"),
    ("costo_usd", "Costo de herramientas imputado al video", "Interna", "Tus suscripciones (fase 7)", "Al publicar"),
    ("horas_produccion", "Horas reales que te llevó el video (todas las etapas)", "Interna", "Tu registro", "Al publicar"),
    ("estado_monetizacion", "Situación en el Programa de Socios y estado de anuncios del video", "Privada", "Studio > Ganar dinero / icono $ del video", "Al publicar y a los 7 d"),
    ("multiplo_7d / 28d", "Vistas ÷ mediana de los 5 videos anteriores del mismo formato", "Calculada", "Fórmula", "7 d y 28 d"),
    ("clasificacion", "Ganador / Normal / Perdedor según los umbrales de Resumen", "Calculada", "Fórmula", "7 d (provisional) y 28 d"),
]


def hoja_glosario(wb):
    ws = wb.create_sheet("Glosario")
    cabecera(ws, 1, ["campo", "qué es", "pública o privada", "dónde se obtiene", "cuándo se registra"])
    anchos(ws, [26, 52, 30, 40, 20])
    for i, f in enumerate(GLOSARIO, start=2):
        for j, v in enumerate(f, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = NEGRO
            c.alignment = AJUSTE
    n = len(GLOSARIO) + 3
    ws.cell(row=n, column=1, value=(
        "Pública = visible para cualquiera (también en canales ajenos). Privada = solo la ve el dueño del canal en "
        "YouTube Studio; para canales ajenos ninguna herramienta la tiene, solo estimaciones (fase 1).")).font = Font(
        name=FUENTE, italic=True, color="666666")
    return ws


# ---------------------------------------------------------------- Leeme
def hoja_leeme(wb):
    ws = wb.active
    ws.title = "Leeme"
    anchos(ws, [3, 110])
    lineas = [
        ("Dashboard de seguimiento del canal", TITULO),
        ("Fase 12 del proyecto. Documentación completa: 14-fase12-dashboard.md.", None),
        ("", None),
        ("Qué celdas rellenas", NEGRITA),
        ("• Texto azul = datos que escribes tú. Texto negro = fórmulas (no tocar). Texto verde = fórmulas que leen otra hoja.", None),
        ("• Fondo amarillo = parámetros y datos de Studio que conviene revisar cada semana (hoja Resumen e Ideas).", None),
        ("• Columnas con cabecera marrón (Videos, desde 'n_formato'): calculadas.", None),
        ("", None),
        ("Rutina", NEGRITA),
        ("1. Al publicar: fila del video en 'Videos' (estado, fecha, URL, duración, costo, horas de producción, cumplimiento).", None),
        ("2. A las 24 h: vistas_24h. A los 7 días: vistas_7d y retencion_30s. A los 28 días: el resto de métricas de Studio.", None),
        ("3. Cada experimento A/B: fila en 'Experimentos' con la hipótesis ANTES de empezar; el resultado al terminar.", None),
        ("4. Cada semana: actualiza los datos amarillos de 'Resumen' con la pantalla 'Ganar dinero' de Studio.", None),
        ("5. Cada mes: revisa 'Ideas' (vuelve a puntuar costo y demanda) y anota aprendizajes en la columna 'aprendizaje'.", None),
        ("", None),
        ("Cómo se identifica un ganador", NEGRITA),
        ("• Múltiplo = vistas a 28 días ÷ mediana de los 5 videos anteriores del mismo formato (largo o Short). Mismo método de la fase 1.", None),
        ("• Ganador: múltiplo ≥ 1,5× y % visto ≥ mediana del formato. Perdedor: múltiplo ≤ 0,67×. A los 7 días se muestra como provisional.", None),
        ("• Los primeros 5 videos de cada formato no tienen 5 anteriores: el múltiplo usa los que haya (1 a 4) y es poco fiable.", None),
        ("• Un ganador se estudia (tema, título, miniatura, gancho) y su rasgo se convierte en hipótesis para el siguiente experimento.", None),
        ("", None),
        ("Ejemplo de fila en 'Videos' (valores ficticios, solo para mostrar el formato; no está en la hoja)", NEGRITA),
        ("id L1 | publicado | 2026-11-19 | Largo | P1 | La verdadera historia de La Llorona | la verdadera historia de la llorona | "
         "duracion_min 10,2 | vistas_24h 180 | vistas_7d 650 | vistas_28d 1.400 | impresiones_28d 21.000 | ctr_28d 5,1 % (escribe 0,051) | "
         "avd_min 3,6 | retencion_30s 68 % | horas_vistas_28d 84 | subs_ganados_28d 9 | trafico_busqueda 35 % | trafico_recomendaciones 40 % | "
         "trafico_shorts 10 % | costo_usd 20,5 | horas_produccion 36 | sin YPP | sí", None),
        ("", None),
        ("Las filas de L1 a L6 y S1 a S13 ya están creadas con las fechas previstas del plan de 90 días (estado 'planificado').", None),
        ("Requisitos del Programa de Socios usados en 'Resumen': fase 10 (support.google.com/youtube/answer/72851 y /12843009, consulta 27-oct-2026).", None),
    ]
    for i, (txt, fuente) in enumerate(lineas, start=1):
        c = ws.cell(row=i, column=2, value=txt)
        c.font = fuente or NEGRO
        c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = Workbook()
    hoja_leeme(wb)
    hoja_videos(wb)
    hoja_ideas(wb)
    hoja_experimentos(wb)
    hoja_resumen(wb)
    hoja_referencias(wb)
    hoja_glosario(wb)
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if c.font and c.font.name != FUENTE:
                    c.font = Font(name=FUENTE, bold=c.font.bold, italic=c.font.italic,
                                  color=c.font.color, size=c.font.size)
    wb.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    main()
